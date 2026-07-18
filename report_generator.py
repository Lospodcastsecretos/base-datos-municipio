import io
import json
import sqlite3
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import database

def generate_report(norma_id: int) -> io.BytesIO:
    conn = sqlite3.connect(database.DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    # 1. Obtener Metadatos de la Norma
    cursor.execute("SELECT * FROM normativas WHERE id = ?", (norma_id,))
    norma = cursor.fetchone()
    
    if not norma:
        conn.close()
        raise ValueError("Norma no encontrada")
        
    doc = Document()
    
    # Título Principal
    title = doc.add_heading("Informe Legal Municipal", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtítulo (Tipo y Número)
    tipo_str = norma['tipo_nombre'] or 'Documento'
    num_str = norma['numero'] or 'S/N'
    subtitle = doc.add_paragraph(f"{tipo_str} Nº {num_str}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.bold = True
    subtitle.runs[0].font.size = Pt(14)
    
    doc.add_paragraph() # Espacio
    
    # Tabla de Metadatos
    doc.add_heading("1. Metadatos Principales", level=1)
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    
    data = [
        ("Título Oficial", norma['titulo'] or "No especificado"),
        ("Categoría / Tema", norma['categoria_nombre'] or "Sin clasificar"),
        ("Fecha de Sanción", norma['fecha'] or "No especificada"),
        ("Estado de Vigencia", "Vigente" if norma['vigente'] else "No Vigente / Derogada"),
        ("Resumen IA", norma['resumen_ia'] or "No generado")
    ]
    
    for i, (key, value) in enumerate(data):
        row_cells = table.rows[i].cells
        row_cells[0].text = key
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = str(value)
        
    doc.add_paragraph()
    
    # Relaciones Jurídicas
    doc.add_heading("2. Relaciones y Afectaciones", level=1)
    salientes = []
    if norma['relaciones_juridicas']:
        try:
            salientes = json.loads(norma['relaciones_juridicas'])
        except:
            pass
            
    if salientes:
        doc.add_paragraph("Esta norma afecta a las siguientes normativas:")
        for rel in salientes:
            dest = rel.get('norma_destino', 'N/A')
            accion = rel.get('accion', 'afecta').upper()
            desc = rel.get('detalle', '')
            bullet = doc.add_paragraph(f"{accion} a la Norma Nº {dest}", style='List Bullet')
            if desc:
                bullet.add_run(f" - {desc}").italic = True
    else:
        doc.add_paragraph("No se registraron afectaciones salientes a otras normativas.")
        
    # Obtener relaciones entrantes
    # Necesitamos buscar en la BD todas las normas que mencionan a esta en su JSON
    # Una forma simple en SQLite es usar LIKE
    search_str = f'%"{num_str}"%'
    cursor.execute("SELECT numero, tipo_nombre, relaciones_juridicas FROM normativas WHERE relaciones_juridicas LIKE ?", (search_str,))
    entrantes_rows = cursor.fetchall()
    entrantes = []
    for r in entrantes_rows:
        try:
            rels = json.loads(r['relaciones_juridicas'])
            for rel in rels:
                if rel.get('norma_destino') == str(num_str):
                    entrantes.append({
                        'origen': f"{r['tipo_nombre']} {r['numero']}",
                        'accion': rel.get('accion', 'afecta').upper(),
                        'detalle': rel.get('detalle', '')
                    })
        except:
            pass
            
    if entrantes:
        doc.add_paragraph("Esta norma recibe afectaciones de:")
        for ent in entrantes:
            bullet = doc.add_paragraph(f"La {ent['origen']} la {ent['accion']}", style='List Bullet')
            if ent['detalle']:
                bullet.add_run(f" - {ent['detalle']}").italic = True
    else:
        doc.add_paragraph("No se registraron modificaciones recibidas de otras normativas.")
        
    doc.add_page_break()
    
    # Articulado y Estructura
    doc.add_heading("3. Estructura y Articulado", level=1)
    
    articulos = database.get_articulos_por_norma(norma_id)
    if articulos:
        doc.add_paragraph(f"Se estructuraron {len(articulos)} artículos en esta normativa:")
        for art in articulos:
            p_art = doc.add_paragraph()
            p_art.add_run(f"Artículo {art['numero']}").bold = True
            
            estado = ""
            if art.get('es_vigente') == 1:
                estado = f" (Vigente, v{art.get('version_numero', 1)})"
            else:
                estado = f" (Derogado el {art.get('fecha_hasta', 'N/A')})"
                
            p_art.add_run(estado).italic = True
            
            doc.add_paragraph(art['texto'])
    else:
        doc.add_paragraph("El documento no se encuentra segmentado en artículos estructurados. A continuación se presenta el texto original extraído:")
        doc.add_paragraph(norma['texto_completo'] or "Texto no disponible.")

    # Guardar a buffer de memoria
    conn.close()
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return file_stream

def generate_rag_report(pregunta: str, respuesta: str, fuentes: list) -> io.BytesIO:
    doc = Document()
    
    title = doc.add_heading("Informe de Consulta - Asistente Jurídico", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    doc.add_heading("1. Consulta Realizada", level=1)
    doc.add_paragraph(pregunta)
    
    doc.add_heading("2. Respuesta de la IA", level=1)
    for par in respuesta.split('\n'):
        if par.strip():
            doc.add_paragraph(par.strip())
            
    doc.add_heading("3. Fuentes Consultadas", level=1)
    if fuentes:
        for f in fuentes:
            doc.add_paragraph(f"{f.get('tipo', 'Doc')} Nº {f.get('numero', 'S/N')} - {f.get('titulo', 'Sin título')}", style='List Bullet')
    else:
        doc.add_paragraph("No se citaron fuentes específicas.")
        
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return file_stream
