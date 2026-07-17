import os
import pdfplumber
import docx
import fitz  # PyMuPDF
import base64
import json
from openai import OpenAI
import google.generativeai as genai
from dotenv import load_dotenv
from spire.doc import Document as SpireDocument

load_dotenv()

def process_document(file_path: str, ia_engine: str = "DeepSeek") -> str:
    """
    Extracts text from a given document file (PDF, DOCX, DOC, JPG, PNG).
    Returns the extracted text as a string.
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return _process_pdf(file_path, ia_engine)
    elif ext == '.docx':
        return _process_docx(file_path)
    elif ext == '.doc':
        return _process_doc(file_path)
    elif ext in ['.jpg', '.jpeg', '.png']:
        return _process_image(file_path, ia_engine)
    else:
        raise ValueError(f"Formato no soportado: {ext}")

def _process_pdf(file_path: str, ia_engine: str) -> str:
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print(f"Error leyendo PDF con pdfplumber {file_path}: {e}")
        
    text = text.strip()
    
    # Si el PDF está vacío o tiene muy poco texto, probablemente es un escaneo (foto)
    if len(text) < 50:
        print(f"PDF {file_path} parece ser un escaneo. Ejecutando OCR con IA...")
        return _ocr_pdf_with_vision(file_path, ia_engine)
        
    return text

def _ocr_pdf_with_vision(file_path: str, ia_engine: str) -> str:
    # Convertir primera(s) pagina(s) a imagen usando PyMuPDF
    ocr_text = ""
    try:
        pdf_document = fitz.open(file_path)
        # Para no exceder límites, extraemos hasta las primeras 5 páginas si es un escaneo
        num_pages = min(5, len(pdf_document)) 
        
        for i in range(num_pages):
            page = pdf_document.load_page(i)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2)) # Buena resolucion
            img_bytes = pix.tobytes("png")
            
            # Pasar los bytes a base64
            img_b64 = base64.b64encode(img_bytes).decode("utf-8")
            
            # Llamar a la API de Vision
            ocr_text += _call_vision_api(img_b64, ia_engine) + "\n\n"
            
        pdf_document.close()
    except Exception as e:
        print(f"Error realizando OCR al PDF: {e}")
        
    return ocr_text.strip()

def _process_image(file_path: str, ia_engine: str) -> str:
    try:
        with open(file_path, "rb") as image_file:
            img_b64 = base64.b64encode(image_file.read()).decode("utf-8")
        return _call_vision_api(img_b64, ia_engine)
    except Exception as e:
        print(f"Error procesando imagen: {e}")
        return ""

def _call_vision_api(img_b64: str, ia_engine: str) -> str:
    prompt = "Extrae TODO el texto visible en esta imagen, manteniendo los saltos de línea. No agregues descripciones, solo el texto exacto que lees."
    
    # Si eligió OpenAI o DeepSeek (que no tiene vision aún), usamos OpenAI si está disponible.
    # Preferimos OpenAI para vision si está configurado, porque Gemini a veces tiene filtros estrictos,
    # pero usamos el que el usuario haya configurado.
    use_openai = False
    
    if "OpenAI" in ia_engine and os.getenv("OPENAI_API_KEY"):
        use_openai = True
    elif "DeepSeek" in ia_engine:
        # DeepSeek no soporta vision. Intentamos usar OpenAI como fallback, si no Gemini
        if os.getenv("OPENAI_API_KEY"):
            use_openai = True
    
    if use_openai and os.getenv("OPENAI_API_KEY"):
        try:
            client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": prompt},
                            {
                                "type": "image_url",
                                "image_url": {"url": f"data:image/png;base64,{img_b64}"}
                            }
                        ]
                    }
                ],
                max_tokens=3000
            )
            return response.choices[0].message.content
        except Exception as e:
            print(f"Error OCR OpenAI: {e}")
            return ""
    else:
        # Fallback a Gemini
        try:
            genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))
            model = genai.GenerativeModel('gemini-1.5-flash')
            # Gemini requiere los bytes crudos y el mime type para la API de dict
            image_parts = [
                {
                    "mime_type": "image/png",
                    "data": base64.b64decode(img_b64)
                }
            ]
            response = model.generate_content([prompt, image_parts[0]])
            return response.text
        except Exception as e:
            print(f"Error OCR Gemini: {e}")
            return ""


def _process_docx(file_path: str) -> str:
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        print(f"Error leyendo DOCX {file_path}: {e}")
    return text.strip()

def _process_doc(file_path: str) -> str:
    text = ""
    try:
        document = SpireDocument()
        document.LoadFromFile(file_path)
        text = document.GetText()
        document.Close()
        
        # Eliminar el watermark gratuito de Spire.Doc si existe
        if "Evaluation Warning: The document was created with Spire.Doc" in text:
            text = text.replace("Evaluation Warning: The document was created with Spire.Doc for Python.", "")
            
    except Exception as e:
        print(f"Error leyendo DOC con Spire: {e}")
        
    return text.strip()
